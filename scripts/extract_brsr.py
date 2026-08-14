import openpyxl, json, re

SRC = '/root/.claude/uploads/ee925557-bbf8-516a-bf50-6db4137ec292/ad345ba6-Green_mentor.xlsx'
wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb['BRSR Reports']

def read_topic_block(col_start, row_start, header_offset=0):
    rows = []
    r = row_start
    while True:
        topic = ws.cell(row=r, column=col_start).value
        if topic is None:
            break
        rows.append({
            "topic": topic,
            "tier": ws.cell(row=r, column=col_start + 1).value,
            "companies": ws.cell(row=r, column=col_start + 2).value,
            "prevalence": ws.cell(row=r, column=col_start + 3).value,
            "mentions": ws.cell(row=r, column=col_start + 4).value,
            "regulatory": ws.cell(row=r, column=col_start + 5).value,
            "opportunity": ws.cell(row=r, column=col_start + 6).value,
            "regulatoryAndOpportunity": ws.cell(row=r, column=col_start + 7).value,
        })
        r += 1
    return rows

topics = {
    "Environment": read_topic_block(1, 16),   # A16..
    "Social": read_topic_block(10, 3),        # J3..
    "Governance": read_topic_block(10, 23),   # J23..
    "Cross-cutting": read_topic_block(19, 3), # S3..
}

def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).replace(",", "").replace("₹", "").strip()
    try:
        return float(s)
    except ValueError:
        return v

largest_emitters = []
r = 21
while ws.cell(row=r, column=19).value is not None:
    largest_emitters.append({
        "company": ws.cell(row=r, column=19).value,
        "fy": ws.cell(row=r, column=20).value,
        "scope1_tCO2e": num(ws.cell(row=r, column=21).value),
        "scope2_tCO2e": num(ws.cell(row=r, column=22).value),
    })
    r += 1

renewable_share = []
r = 38
while ws.cell(row=r, column=1).value is not None:
    renewable_share.append({
        "company": ws.cell(row=r, column=1).value,
        "fy": ws.cell(row=r, column=2).value,
        "renewablePct": ws.cell(row=r, column=3).value,
    })
    r += 1

ltifr_distribution = []
r = 38
while ws.cell(row=r, column=5).value is not None:
    ltifr_distribution.append({
        "band": ws.cell(row=r, column=5).value,
        "companies": ws.cell(row=r, column=6).value,
    })
    r += 1

largest_water_withdrawers = []
r = 38
while ws.cell(row=r, column=8).value is not None:
    largest_water_withdrawers.append({
        "company": ws.cell(row=r, column=8).value,
        "fy": ws.cell(row=r, column=9).value,
        "withdrawalKL": num(ws.cell(row=r, column=10).value),
    })
    r += 1

companies = []
r = 39
while True:
    symbol = ws.cell(row=r, column=19).value
    if symbol is None:
        break
    sector_raw = ws.cell(row=r, column=22).value
    nic_section = None
    nic_section_name = None
    if sector_raw and sector_raw != "unmapped":
        m = re.match(r"^([A-Z])(.*)$", sector_raw)
        if m:
            nic_section, nic_section_name = m.group(1), m.group(2)
    companies.append({
        "symbol": symbol,
        "name": ws.cell(row=r, column=20).value,
        "fy": ws.cell(row=r, column=21).value,
        "sectorRaw": sector_raw,
        "nicSection": nic_section,
        "nicSectionName": nic_section_name,
        "coverage": ws.cell(row=r, column=23).value,
    })
    r += 1

out = {
    "topics": topics,
    "largestEmitters": largest_emitters,
    "renewableShare": renewable_share,
    "ltifrDistribution": ltifr_distribution,
    "largestWaterWithdrawers": largest_water_withdrawers,
    "companies": companies,
}

with open('/home/user/green-mentor/src/data/brsr.json', 'w') as f:
    json.dump(out, f, indent=2)

print("topics per pillar:", {k: len(v) for k, v in topics.items()})
print("companies:", len(companies))
print("unmapped sectors:", sum(1 for c in companies if c["sectorRaw"] == "unmapped"))
print("sample company:", companies[0])
